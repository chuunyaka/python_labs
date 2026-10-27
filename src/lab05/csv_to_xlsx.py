import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использовать openpyxl.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    c_path = Path(csv_path)
    x_path = Path(xlsx_path)

    if not c_path.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")
    
    with c_path.open("r", newline="", encoding="utf-8-sig") as f:
        first_line = f.readline()
        if not first_line.strip():
            raise ValueError("CSV файл пустой")

        if ";" in first_line:
            delimiter = ";"
        else:
            delimiter = ","

        f.seek(0)
        reader = csv.reader(f,delimiter=delimiter)
        rows = list(reader)

    if not rows or not any(rows[0]):
        raise ValueError("CSV файл пуст или некорректен")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    for row in rows:
        ws.append(row)
        
    for col_idx, col in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        
        max_len = 0
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 5, 8)
        
    x_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(x_path)